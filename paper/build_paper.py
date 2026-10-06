"""Generate the final ICRE 2026 paper from measured project results."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "paper" / "template_icre_2026.docx"
METRICS = ROOT / "results" / "metrics.json"
OUTPUT = ROOT / "paper" / "ICRE_2026_Final_Bilingual_Sentiment_Paper.docx"

TITLE = "Bilingual Arabic-English Customer Review Sentiment Analysis Using Machine Learning"
AUTHORS = "Mohammad Rouhi Khattab, Salam Al-Buhaisi, and Mohammad Abu Khammash"


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def set_font(run, name="Times New Roman", size=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_container(container):
    element = container._element
    for child in list(element):
        element.remove(child)
    return container.add_paragraph()


def set_columns(section, count, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def configure_section(section, columns):
    section.page_width = Inches(8.267)
    section.page_height = Inches(11.694)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)
    section.gutter = 0
    set_columns(section, columns)


def set_header_footer(doc):
    line = (
        "Proc. of the 9th International Conference on Resilient Engineering "
        "(ICRE 2026), 8-10 December 2026, Gaza, Palestine"
    )
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for header in (section.header, section.first_page_header):
            header.is_linked_to_previous = False
            p = clear_container(header)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            set_font(r, size=9, italic=True)
        for footer in (section.footer, section.first_page_footer):
            footer.is_linked_to_previous = False
            p = clear_container(footer)
            p.paragraph_format.space_after = Pt(0)


def add_labeled_paragraph(doc, style, label, text):
    p = doc.add_paragraph(style=style)
    a = p.add_run(label)
    b = p.add_run(text)
    if style == "Keywords":
        a.bold = True
        a.italic = True
        b.bold = False
        b.italic = True
    return p


def add_body(doc, text):
    return doc.add_paragraph(text, style="Body Text")


def set_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    for run in paragraph.runs:
        rpr = run._element.get_or_add_rPr()
        rtl = rpr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            rpr.append(rtl)
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:bidi"), "ar-SA")
        set_font(run, size=9)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_twips):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "808080")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side in ("top", "left", "bottom", "right"):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), "70")
                node.set(qn("w:type"), "dxa")


def add_compact_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, size=7.5, bold=True)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_font(run, size=7.5)
    set_table_geometry(table, widths)
    return table


def add_figure(doc, path, caption, width=3.10, alt_text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph(caption, style="figure caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    metrics = json.loads(METRICS.read_text(encoding="utf-8"))
    selected = metrics["selected_model"]
    selected_metrics = metrics["models"][selected]
    lr = metrics["models"]["Logistic Regression"]
    nb = metrics["models"]["Multinomial Naive Bayes"]
    dataset = metrics["dataset"]
    total = dataset["total_records"]
    english = dataset["language_counts"]["en"]
    arabic = dataset["language_counts"]["ar"]

    abstract = (
        f"This paper presents a reproducible Arabic-English sentiment analysis prototype for "
        f"customer-review support. After deduplication, the dataset contains {total:,} records: "
        f"{english:,} English review sentences from the UCI Sentiment Labelled Sentences "
        f"collection and {arabic:,} non-zero Arabic entries from a translated AFINN lexicon. "
        f"Unicode-aware normalization and character 3-5-gram TF-IDF features were used to "
        f"compare Logistic Regression with Multinomial Naive Bayes. On a stratified 20% "
        f"hold-out set, {selected} achieved {pct(selected_metrics['test']['accuracy'])} "
        f"accuracy and {pct(selected_metrics['test']['f1_macro'])} macro F1. Accuracy was "
        f"{pct(selected_metrics['by_language']['en']['accuracy'])} on English review "
        f"sentences and {pct(selected_metrics['by_language']['ar']['accuracy'])} on held-out "
        f"Arabic lexicon entries. The delivered system includes an offline bilingual web "
        f"interface, command-line prediction, reproducible training, tests, and saved results."
    )
    arabic_summary = (
        f"يعرض هذا البحث نظاماً عملياً لتحليل المشاعر في المراجعات العربية والإنجليزية. "
        f"تتكون البيانات بعد إزالة التكرار من {total:,} سجلاً، منها {english:,} جملة مراجعة "
        f"إنجليزية حقيقية و{arabic:,} مفردة أو عبارة عربية من معجم AFINN المترجم. استُخدمت "
        f"خصائص TF-IDF على المقاطع الحرفية، وتمت مقارنة الانحدار اللوجستي مع نايف بايز. حقق "
        f"النموذج المختار دقة كلية مقدارها {pct(selected_metrics['test']['accuracy'])}، مع "
        f"التأكيد أن التقييم العربي معجمي ولا يمثل بعدُ تقييماً على مجموعة كبيرة من مراجعات "
        f"عربية طبيعية."
    )

    doc = Document(TEMPLATE)
    clear_body(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Final AI course project paper formatted for ICRE 2026"
    doc.core_properties.author = AUTHORS
    doc.core_properties.keywords = (
        "sentiment analysis; Arabic NLP; bilingual classification; TF-IDF; machine learning"
    )

    title = doc.add_paragraph(style="paper title")
    title.add_run(TITLE)
    authors = doc.add_paragraph(style="Author")
    authors.paragraph_format.space_before = Pt(7)
    authors.paragraph_format.space_after = Pt(10)
    authors.add_run(AUTHORS)

    doc.add_section(WD_SECTION_START.CONTINUOUS)
    configure_section(doc.sections[0], 1)
    configure_section(doc.sections[1], 2)
    set_header_footer(doc)

    add_labeled_paragraph(doc, "Abstract", "Abstract- ", abstract)
    add_labeled_paragraph(
        doc,
        "Keywords",
        "Keywords- ",
        "sentiment analysis, Arabic NLP, bilingual classification, character n-grams, TF-IDF, machine learning",
    )
    arabic_head = doc.add_paragraph("Arabic Summary", style="Heading 5")
    arabic_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    arabic_paragraph = add_body(doc, arabic_summary)
    set_rtl(arabic_paragraph)

    doc.add_paragraph("Introduction", style="Heading 1")
    add_body(
        doc,
        "Online reviews provide direct evidence about product quality, service failures, and customer satisfaction, but manual reading becomes impractical when review volume grows. Sentiment analysis addresses this problem by assigning polarity to text and is a widely studied application of natural language processing (NLP) [1]. A practical classifier can help an organization screen feedback, identify dissatisfied customers, and prioritize recurring problems without requiring personal identifiers.",
    )
    add_body(
        doc,
        "Arabic-English support is especially relevant in multilingual environments where customers may switch between Arabic, English, or mixed text. Arabic introduces additional challenges, including optional diacritics, multiple forms of the same letter, rich morphology, dialectal spelling, and right-to-left presentation. Large transformer models can address some of these issues, but they require more data and computing resources. This project therefore investigates whether a compact, fully offline machine-learning pipeline can provide an understandable baseline suitable for an undergraduate AI project.",
    )
    add_body(
        doc,
        "The contributions are fourfold. First, a reproducible bilingual dataset is assembled from clearly attributed public resources. Second, one character-level TF-IDF representation is used for both scripts. Third, two classic classifiers are compared using a fixed hold-out set and five-fold cross-validation. Fourth, the selected model is deployed through bilingual command-line and local web interfaces. All reported values were generated by the delivered code, and the limitations of the Arabic component are reported explicitly.",
    )

    doc.add_paragraph("Related Work", style="Heading 1")
    add_body(
        doc,
        "Pang and Lee [1] describe sentiment analysis as the computational study of opinions and evaluations in text. Traditional supervised systems commonly transform text into sparse features and train linear or probabilistic classifiers. The UCI Sentiment Labelled Sentences dataset was introduced in connection with the work of Kotzias et al. [2] and provides balanced positive and negative sentences from Amazon, IMDb, and Yelp [3]. Its small size and multiple domains make it useful for a transparent classroom baseline.",
    )
    add_body(
        doc,
        "Lexicon-based approaches use lists of terms with manually assigned sentiment scores. AFINN assigns integer valence values from negative to positive terms and was evaluated for short social-media text [4]. The multilingual software source used in this project distributes translated AFINN lists, including Arabic [5]. Lexicons are computationally efficient, but isolated words cannot represent the complete pragmatics of review sentences, especially negation, sarcasm, and dialect. The Arabic results in this paper must therefore be read as lexicon-term generalization rather than full review-corpus validation.",
    )
    add_body(
        doc,
        "Modern contextual models such as BERT learn bidirectional representations from large unlabeled corpora and can be fine-tuned for classification [7]. They are a strong future direction, but the present work deliberately uses scikit-learn [6] and interpretable sparse features to keep training reproducible on an ordinary computer without cloud services or specialized hardware.",
    )

    doc.add_paragraph("Dataset and Ethics", style="Heading 1")
    add_body(
        doc,
        f"The preprocessing script reads the three labelled UCI sources and the Arabic AFINN file, removes empty records, normalizes text, and removes exact duplicates within each language. The final dataset contains {total:,} records. The English portion contains {english:,} review sentences after deduplication. The Arabic portion contains {arabic:,} non-zero terms or short phrases. Table I summarizes the data used for training and evaluation.",
    )
    cap = doc.add_paragraph("DATASET COMPOSITION AFTER DEDUPLICATION", style="table head")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_compact_table(
        doc,
        ["Component", "Records", "Unit"],
        [
            ("UCI Amazon", dataset["source_counts"]["UCI-amazon"], "Review sentence"),
            ("UCI IMDb", dataset["source_counts"]["UCI-imdb"], "Review sentence"),
            ("UCI Yelp", dataset["source_counts"]["UCI-yelp"], "Review sentence"),
            ("AFINN Arabic", dataset["source_counts"]["AFINN-ar-translated-lexicon"], "Term / phrase"),
        ],
        [1500, 800, 2500],
    )
    add_body(
        doc,
        "No names, account identifiers, phone numbers, addresses, or health records are required by the system. New institutional data should be anonymized before use. The UCI resource is published under CC BY 4.0, while the multilingual AFINN software source is distributed under the MIT license. The project retains source labels so that performance can be audited separately rather than hidden inside one overall score.",
    )

    doc.add_paragraph("Methodology", style="Heading 1")
    doc.add_paragraph("Text Normalization and Representation", style="Heading 2")
    add_body(
        doc,
        "Text is normalized with Unicode NFKC, lowercasing for Latin text, removal of Arabic diacritics and tatweel, normalization of common alif variants, URL and e-mail replacement, punctuation filtering, and whitespace collapse. The same deterministic function is used during data preparation and inference. This avoids training-serving mismatch and preserves both Arabic and English characters.",
    )
    add_body(
        doc,
        f"Each record is represented using character n-gram TF-IDF with n from 3 to 5, sublinear term frequency, L2 normalization, a minimum document frequency of two, and a maximum of 60,000 features. Character features reduce dependence on language-specific tokenizers and capture useful subword patterns. The selected fitted pipeline retained {metrics['feature_extraction']['feature_count_selected_model']:,} features.",
    )
    doc.add_paragraph("Models and Evaluation", style="Heading 2")
    add_body(
        doc,
        f"Two classifiers were compared: L2-regularized Logistic Regression with balanced class weights and Multinomial Naive Bayes with additive smoothing. A stratified split retained {metrics['split']['train_records']:,} records for training and {metrics['split']['test_records']:,} for testing while preserving every language-label combination. Five-fold stratified cross-validation was also run on the complete dataset. Accuracy, macro precision, macro recall, and macro F1 are reported. Model selection used hold-out macro F1 because the combined dataset is not perfectly class-balanced.",
    )
    add_body(
        doc,
        "The learned task is binary: positive versus negative. The user interface additionally returns neutral/uncertain when the maximum predicted probability is below 0.60. This is an abstention mechanism and is not reported as a learned third class. It prevents weak predictions from being presented as confident decisions, but it should not be confused with a true neutral-sentiment model.",
    )

    doc.add_paragraph("System Implementation", style="Heading 1")
    add_body(
        doc,
        "The implementation is written in Python and uses pandas for data handling, scikit-learn for feature extraction and classifiers, joblib for model persistence, and matplotlib for evaluation figures. The training script creates a processed UTF-8 CSV file, compares both pipelines, writes structured JSON metrics, saves all test predictions, selects the best model, and exports charts. A unit-test suite checks Arabic normalization, language detection, and positive/negative predictions in both languages.",
    )
    add_body(
        doc,
        "Two interfaces are supplied. The command-line program accepts a review and can return either readable bilingual output or JSON. The local web application uses Python's standard HTTP server and therefore adds no web-framework dependency. It runs on localhost, accepts Arabic, English, and mixed text, and shows the bilingual label and confidence. Because the application is offline, review text is not transmitted to an external service.",
    )

    doc.add_paragraph("Results and Discussion", style="Heading 1")
    add_body(
        doc,
        f"Logistic Regression produced the highest hold-out macro F1 and was selected for delivery. It achieved {pct(lr['test']['accuracy'])} accuracy and {pct(lr['test']['f1_macro'])} macro F1. Its five-fold mean accuracy was {pct(lr['cross_validation']['accuracy_mean'])} with a standard deviation of {lr['cross_validation']['accuracy_std'] * 100:.2f} percentage points. Multinomial Naive Bayes reached {pct(nb['test']['accuracy'])} accuracy and {pct(nb['test']['f1_macro'])} macro F1. The close cross-validation and hold-out values suggest that the comparison is stable under the fixed experimental design.",
    )
    cap = doc.add_paragraph("MODEL PERFORMANCE ON THE HOLD-OUT SET", style="table head")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_compact_table(
        doc,
        ["Model", "Acc.", "Macro F1", "CV F1"],
        [
            (
                "Logistic Regression",
                pct(lr["test"]["accuracy"]),
                pct(lr["test"]["f1_macro"]),
                pct(lr["cross_validation"]["f1_macro_mean"]),
            ),
            (
                "Multinomial NB",
                pct(nb["test"]["accuracy"]),
                pct(nb["test"]["f1_macro"]),
                pct(nb["cross_validation"]["f1_macro_mean"]),
            ),
        ],
        [1700, 900, 1100, 1100],
    )
    add_figure(
        doc,
        ROOT / "results" / "model_comparison.png",
        "Hold-out accuracy and macro F1 for both candidate models.",
        alt_text="Bar chart comparing hold-out accuracy and macro F1 for Logistic Regression and Multinomial Naive Bayes.",
    )
    add_body(
        doc,
        f"The selected model performed substantially better on English review sentences than on the Arabic lexicon component. English accuracy was {pct(lr['by_language']['en']['accuracy'])} with {pct(lr['by_language']['en']['f1_macro'])} macro F1. Arabic lexicon accuracy was {pct(lr['by_language']['ar']['accuracy'])} with {pct(lr['by_language']['ar']['f1_macro'])} macro F1. Among English sources, accuracy was {pct(lr['by_source']['UCI-amazon']['accuracy'])} for Amazon, {pct(lr['by_source']['UCI-imdb']['accuracy'])} for IMDb, and {pct(lr['by_source']['UCI-yelp']['accuracy'])} for Yelp.",
    )
    add_figure(
        doc,
        ROOT / "results" / "confusion_matrix.png",
        "Confusion matrix for the selected Logistic Regression model.",
        alt_text="Confusion matrix showing negative and positive predictions for the selected Logistic Regression model.",
    )
    add_body(
        doc,
        f"The confusion matrix contains {lr['confusion_matrix'][0][0]} correctly predicted negative records and {lr['confusion_matrix'][1][1]} correctly predicted positive records. There were {lr['confusion_matrix'][0][1]} false positives and {lr['confusion_matrix'][1][0]} false negatives. Errors were therefore distributed across both classes rather than being confined to one label. The balanced class weights helped Logistic Regression despite the larger number of negative Arabic lexicon entries.",
    )
    add_body(
        doc,
        "The results show that a compact character-level baseline can support both scripts with one feature pipeline. However, the Arabic and English numbers are not directly equivalent: English test records are natural review sentences, while Arabic test records are translated sentiment terms. The lower Arabic score and this unit mismatch both motivate a dedicated Arabic review corpus rather than stronger claims based on the overall result alone.",
    )

    doc.add_paragraph("Limitations and Future Work", style="Heading 1")
    add_body(
        doc,
        "The main limitation is the absence of a large, naturally occurring, manually annotated Arabic customer-review dataset. The translated lexicon does not capture sentence-level negation, sarcasm, word order, Palestinian dialect, Arabizi, or code-switching. The English dataset is also small and limited to three older domains. Random sentence splitting measures in-domain generalization but does not prove transfer to a new organization or time period.",
    )
    add_body(
        doc,
        "The project also uses binary labels. Its neutral output is only a confidence-based abstention, and probability estimates from simple text classifiers should not be interpreted as calibrated risk without additional validation. Customer-service deployment would require human review, drift monitoring, domain-specific thresholds, and an appeal path for decisions that affect users.",
    )
    add_body(
        doc,
        "Future work should collect consented, anonymized Arabic and mixed-language reviews from a relevant local domain, define annotation guidelines, measure inter-annotator agreement, and keep a source-separated test set. Word and character features can then be compared with multilingual BERT or an Arabic pretrained model [7]. A genuine three-class dataset should be used if neutral sentiment is required. Explainability can also be improved by highlighting influential n-grams and grouping errors by dialect, topic, and review length.",
    )

    doc.add_paragraph("Conclusion", style="Heading 1")
    add_body(
        doc,
        f"A complete Arabic-English sentiment-analysis prototype was implemented, tested, and documented. The system combines Unicode normalization, character n-gram TF-IDF, and two classic classifiers in a reproducible pipeline. Logistic Regression was selected with {pct(lr['test']['accuracy'])} overall accuracy, while the English review subset reached {pct(lr['by_language']['en']['accuracy'])}. The system is suitable as an educational baseline and offline demonstration, but the Arabic lexicon evaluation is not sufficient for claims about real Arabic customer reviews. The delivered code, model, predictions, metrics, figures, interfaces, and paper provide a transparent foundation for collecting stronger local data and evaluating modern multilingual models.",
    )

    refs_head = doc.add_paragraph("References", style="Heading 5")
    refs_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    references = [
        "B. Pang and L. Lee, “Opinion mining and sentiment analysis,” Foundations and Trends in Information Retrieval, vol. 2, nos. 1-2, pp. 1-135, 2008.",
        "D. Kotzias, M. Denil, N. de Freitas, and P. Smyth, “From group to individual labels using deep features,” in Proc. 21st ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining, 2015, pp. 597-606.",
        "D. Kotzias, “Sentiment Labelled Sentences,” UCI Machine Learning Repository, 2015, doi: 10.24432/C57604.",
        "F. A. Nielsen, “A new ANEW: Evaluation of a word list for sentiment analysis in microblogs,” in Proc. ESWC Workshop on Making Sense of Microposts, vol. 718, 2011, pp. 93-98.",
        "M. Barile, “multilang-sentiment: Multi language AFINN-based sentiment analysis,” GitHub repository, 2021. [Online]. Available: https://github.com/marcellobarile/multilang-sentiment",
        "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
        "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, “BERT: Pre-training of deep bidirectional transformers for language understanding,” in Proc. NAACL-HLT, 2019, pp. 4171-4186.",
    ]
    for ref in references:
        doc.add_paragraph(ref, style="references")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
